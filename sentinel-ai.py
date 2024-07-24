import requests
from PyPDF2 import PdfReader
import io
import re
from docx import Document
import openpyxl
import csv
from bs4 import BeautifulSoup
import json
from colorama import init, Fore, Style
import tkinter as tk
from tkinter import filedialog, messagebox
import openai

# Set your OpenAI API key here
openai.api_key = 'YOUR OPEN AI API KEY'

# Initialize Colorama
init(autoreset=True)

def display_banner():
    print(Fore.CYAN + Style.BRIGHT + "\n" + "="*60)
    print(Fore.CYAN + Style.BRIGHT + " " * 10 + "Sensitive Data Scanner - Sentinel AI")
    print(Fore.CYAN + Style.BRIGHT + " " * 18 + "Creator: Alperen Ugurlu")
    print(Fore.CYAN + Style.BRIGHT + "="*60 + "\n" + Style.RESET_ALL)

def load_file(file_path):
    file_type = file_path.split('.')[-1]
    return file_path, file_type

def extract_text_from_pdf(pdf_path):
    pdf_reader = PdfReader(pdf_path)
    text = ""
    for page_num in range(len(pdf_reader.pages)):
        page = pdf_reader.pages[page_num]
        text += page.extract_text()
    return text

def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text
    return text

def extract_text_from_excel(excel_path):
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    text = ""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text

def extract_text_from_csv(csv_path):
    with open(csv_path, newline='', encoding='utf-8') as csv_file:
        reader = csv.reader(csv_file)
        text = "\n".join([" ".join(row) for row in reader])
    return text

def extract_text_from_html(html_path):
    with open(html_path, 'r', encoding='utf-8') as html_file:
        soup = BeautifulSoup(html_file, 'html.parser')
        text = soup.get_text()
    return text

def extract_text_from_json(json_path):
    with open(json_path, 'r', encoding='utf-8') as json_file:
        data = json.load(json_file)
        text = json.dumps(data, indent=4)
    return text

def find_sensitive_data(text, patterns):
    found_data = {key: re.findall(pattern, text) for key, pattern in patterns.items()}
    return found_data

def find_sensitive_data_with_ai(text, model="gpt-3.5-turbo"):
    response = openai.ChatCompletion.create(
        model=model,
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": f"Find all sensitive information such as names, organizations, locations, and any other personal or sensitive data in the following text:\n\n{text}\n\nList them under the categories: names, organizations, locations, miscellaneous."}
        ],
        max_tokens=1000
    )
    response_text = response['choices'][0]['message']['content'].strip()
    sensitive_data = {"names": [], "organizations": [], "locations": [], "miscellaneous": []}

    lines = response_text.split('\n')
    current_category = None

    for line in lines:
        if 'names:' in line.lower():
            current_category = 'names'
        elif 'organizations:' in line.lower():
            current_category = 'organizations'
        elif 'locations:' in line.lower():
            current_category = 'locations'
        elif 'miscellaneous:' in line.lower():
            current_category = 'miscellaneous'
        elif current_category:
            sensitive_data[current_category].append(line.strip('- '))

    return sensitive_data

def print_sensitive_data(file_path, sensitive_data):
    print(Fore.CYAN + f"\nAnalyzing file: {file_path}")
    if any(sensitive_data.values()):
        print(Fore.RED + "Sensitive Data Found:")
        for key, matches in sensitive_data.items():
            if matches:
                print(Fore.YELLOW + f"  {key.capitalize()}:")
                for match in matches:
                    print(Fore.GREEN + f"    - {match}")
    else:
        print(Fore.GREEN + "No sensitive data found.")

def save_results_to_file(file_path, sensitive_data):
    with open("sensitive_data_report.txt", "a") as report_file:
        report_file.write(f"\nAnalyzing file: {file_path}\n")
        if any(sensitive_data.values()):
            report_file.write("Sensitive Data Found:\n")
            for key, matches in sensitive_data.items():
                if matches:
                    report_file.write(f"  {key.capitalize()}:\n")
                    for match in matches:
                        report_file.write(f"    - {match}\n")
        else:
            report_file.write("No sensitive data found.\n")

def process_file(file_path, model):
    try:
        file_path, file_type = load_file(file_path)

        if file_type == 'pdf':
            text = extract_text_from_pdf(file_path)
        elif file_type == 'docx':
            text = extract_text_from_docx(file_path)
        elif file_type == 'xlsx':
            text = extract_text_from_excel(file_path)
        elif file_type == 'csv':
            text = extract_text_from_csv(file_path)
        elif file_type == 'html':
            text = extract_text_from_html(file_path)
        elif file_type == 'json':
            text = extract_text_from_json(file_path)
        else:
            print(Fore.RED + f"Unsupported file type: {file_type}")
            return

        sensitive_data_regex = find_sensitive_data(text, sensitive_patterns)
        sensitive_data_ai = find_sensitive_data_with_ai(text, model=model)
        sensitive_data = {**sensitive_data_regex, **sensitive_data_ai}
        print_sensitive_data(file_path, sensitive_data)
        save_results_to_file(file_path, sensitive_data)

    except Exception as e:
        print(Fore.RED + f"Error processing {file_path}: {e}")

def browse_files(model):
    file_paths = filedialog.askopenfilenames()
    for file_path in file_paths:
        process_file(file_path, model)
    messagebox.showinfo("Process Complete", "File scanning completed.")

def select_gpt35():
    browse_files("gpt-3.5-turbo")

def select_gpt4():
    browse_files("gpt-4")

# Sensitive data patterns
sensitive_patterns = {
    # Identity Information
    'email': r'[a-zA-Z0-9+_.-]+@[a-zA-Z0-9.-]+',
    'phone': r'\b\d{3}[-.]?\d{3}[-.]?\d{4}\b',
    'credit_card': r'\b(?:\d{4}[-.\s]?){3}\d{4}\b',
    'ssn': r'\b\d{3}-\d{2}-\d{4}\b',
    'ip_address': r'\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b',
    'admin': r'\badmin\b',
    'user': r'\buser\b',
    'password': r'\bpassword\b|\bpasswd\b',
    'date_of_birth': r'\b\d{2}/\d{2}/\d{4}\b|\b\d{2}-\d{2}-\d{4}\b',
    'national_identity_number': r'\b\d{11}\b',

    # Financial Information
    'bank_account': r'\b\d{2,3}[-.\s]?\d{6,8}[-.\s]?\d{1,3}\b',
    'credit_card_security_code': r'\b\d{3,4}\b',
    'international_bank_account_number': r'\b[A-Z]{2}\d{2}[A-Z0-9]{1,30}\b',
    'swift_code': r'\b[A-Z]{4}[-.]?[A-Z]{2}[-.]?[A-Z0-9]{2}[-.]?[A-Z0-9]{3}?\b',
    'bitcoin_address': r'\b[13][a-km-zA-HJ-NP-Z1-9]{25,34}\b',

    # Personal Information
    'passport_number': r'\b[A-Z]{1,2}\d{6,7}\b',
    'driver_license': r'\b[A-Z0-9]{1,9}\b',
    'address': r'\b\d{1,5}\s\w+\s\w+(\s\w+){0,2},?\s\w+,\s\w{2}\s\d{5}\b',
    'health_insurance_number': r'\b\d{3}-\d{2}-\d{4}\b',
    'medical_record_number': r'\b\d{2,3}[-.\s]?\d{2,3}[-.\s]?\d{4}\b',
    'social_media_handle': r'@[A-Za-z0-9_]{1,15}',

    # Corporate Information
    'tax_identification_number': r'\b\d{2}-\d{7}\b',
    'business_identification_number': r'\b\d{9}\b',
    'corporate_email': r'[a-zA-Z0-9._%+-]+@(?:[a-zA-Z0-9-]+\.)+[a-zA-Z]{2,}\b',
    'employment_id': r'\b[A-Z0-9]{8,12}\b',

    # Other
    'url': r'https?://(?:[-\w.]|(?:%[\da-fA-F]{2}))+',
    'vehicle_identification_number': r'[A-HJ-NPR-Z0-9]{17}',
    'mac_address': r'([0-9A-Fa-f]{2}[:-]){5}([0-9A-Fa-f]{2})',
    'imei': r'\b\d{15}\b',
    'serial_number': r'\b[A-Z0-9]{10,12}\b',
    'api_key': r'\b[A-Za-z0-9]{32}\b',
    'cryptographic_key': r'\b[A-Fa-f0-9]{64}\b',
    'personal_identification_number': r'\b\d{4,6}\b',
    'financial_account_number': r'\b\d{4}[-.\s]?\d{4}[-.\s]?\d{4}[-.\s]?\d{4}\b',
    'student_id': r'\b[A-Z0-9]{8}\b',
    'library_card_number': r'\b\d{8,10}\b',
    'membership_number': r'\b[A-Z0-9]{6,12}\b',
    'insurance_policy_number': r'\b[A-Z0-9]{10,12}\b',
    'employee_number': r'\b[A-Z0-9]{6,10}\b',
    'security_clearance_number': r'\b[A-Z0-9]{6,8}\b',
    'network_security_key': r'\b[A-Za-z0-9]{8,64}\b',
    'birth_certificate_number': r'\b\d{8,10}\b',
    'census_record': r'\b\d{8,12}\b',
    'patent_number': r'\b\d{7}\b',
    'trademark_number': r'\b\d{7}\b',
    'court_case_number': r'\b[A-Z0-9]{6,10}\b',
    'inmate_number': r'\b\d{6}\b',
    'military_id': r'\b[A-Z0-9]{6,10}\b',
    'voter_registration_number': r'\b[A-Z0-9]{6,10}\b',
    'nonprofit_registration_number': r'\b[A-Z0-9]{6,10}\b',
    'trading_license_number': r'\b[A-Z0-9]{6,10}\b'
}

# Create the GUI
root = tk.Tk()
root.title("Sensitive Data Scanner - Sentinel AI")
root.geometry("500x400")

label = tk.Label(root, text="Select files to scan for sensitive data\nCreator: Alperen Ugurlu")
label.pack(pady=20)

gpt35_button = tk.Button(root, text="Use GPT-3.5", command=select_gpt35)
gpt35_button.pack(pady=10)

gpt4_button = tk.Button(root, text="Use GPT-4", command=select_gpt4)
gpt4_button.pack(pady=10)

# Display the banner when the program starts
display_banner()

root.mainloop()
